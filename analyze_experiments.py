"""
Analyze experiments folder and create Excel priority list
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime

# Analysis of experiments folder
experiments_analysis = [
    # CRITICAL PRIORITY - Multi-Agent Architecture (Most Advanced)
    {
        "Priority": "CRITICAL",
        "File": "11_multi_agent_research_system.py",
        "Category": "Multi-Agent Architecture",
        "Learning Value": "⭐⭐⭐⭐⭐",
        "Key Skills": "Coordinator pattern, agent communication, role-based specialists, task routing",
        "Complexity": "Advanced",
        "Lines": "892",
        "Integration Value": "Immediate - shows production multi-agent pattern for Week 5",
        "Status": "Ready to Clone",
        "Notes": "4 specialized agents (Coordinator, Researcher, Analyst, Writer). Shows message passing, task orchestration, and team coordination."
    },
    
    # CRITICAL PRIORITY - Advanced Memory Systems
    {
        "Priority": "CRITICAL",
        "File": "09_research_assistant_smart_memory.py",
        "Category": "Memory & Context Management",
        "Learning Value": "⭐⭐⭐⭐⭐",
        "Key Skills": "Semantic search, embedding-based memory, context relevance scoring, smart retrieval",
        "Complexity": "Advanced",
        "Lines": "600+",
        "Integration Value": "High - solves context window optimization for long conversations",
        "Status": "Ready to Clone",
        "Notes": "Implements semantic memory with relevance scoring. Critical for scaling conversations beyond token limits."
    },
    
    # CRITICAL PRIORITY - Token Tracking & Cost Management
    {
        "Priority": "CRITICAL",
        "File": "14_research_assistant_token_tracking.py",
        "Category": "Performance & Cost Optimization",
        "Learning Value": "⭐⭐⭐⭐⭐",
        "Key Skills": "Token counting, cost calculation, budget limits, performance metrics",
        "Complexity": "Intermediate",
        "Lines": "400+",
        "Integration Value": "Immediate - needed for Week 4 benchmarking and metrics",
        "Status": "Ready to Clone",
        "Notes": "Tracks tokens per request, cumulative costs, sets budget alerts. Essential for production deployment."
    },
    
    # HIGH PRIORITY - API Integration Patterns
    {
        "Priority": "HIGH",
        "File": "05_research_assistant_with_apis.py",
        "Category": "External Data Integration",
        "Learning Value": "⭐⭐⭐⭐⭐",
        "Key Skills": "Multi-API orchestration, arXiv, Wikipedia, GitHub, caching strategies",
        "Complexity": "Intermediate",
        "Lines": "500+",
        "Integration Value": "High - shows how to integrate 5+ free APIs without auth",
        "Status": "Ready to Clone",
        "Notes": "Zero-config API integration (arXiv, Wikipedia, Dictionary, GitHub, Open Library). No API keys needed!"
    },
    
    # HIGH PRIORITY - Export & Documentation
    {
        "Priority": "HIGH",
        "File": "15_research_assistant_export.py",
        "Category": "Data Export & Reporting",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Multi-format export (MD, PDF, JSON), templating, formatting",
        "Complexity": "Intermediate",
        "Lines": "450+",
        "Integration Value": "Medium - useful for generating reports from research findings",
        "Status": "Ready to Clone",
        "Notes": "Exports research to Markdown, PDF, JSON. Includes templates and formatting."
    },
    
    # HIGH PRIORITY - Knowledge Base Integration
    {
        "Priority": "HIGH",
        "File": "16_research_assistant_obsidian.py",
        "Category": "Knowledge Management",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Obsidian vault integration, auto-tagging, backlinks, graph knowledge",
        "Complexity": "Intermediate",
        "Lines": "500+",
        "Integration Value": "Medium - creates structured knowledge base from research",
        "Status": "Ready to Clone",
        "Notes": "Automatically creates Obsidian notes with tags, backlinks, and relationships. Builds knowledge graph."
    },
    
    # HIGH PRIORITY - PDF Analysis
    {
        "Priority": "HIGH",
        "File": "12_research_assistant_with_pdf.py",
        "Category": "Document Processing",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "PDF parsing, chunk-based analysis, relevance scoring, batch processing",
        "Complexity": "Intermediate",
        "Lines": "400+",
        "Integration Value": "Medium - useful for document analysis workflows",
        "Status": "Ready to Clone",
        "Notes": "Analyzes PDFs in chunks, scores relevance, batch processes directories. Includes cost optimization."
    },
    
    # HIGH PRIORITY - Visual Diagrams
    {
        "Priority": "HIGH",
        "File": "13_research_assistant_diagrams.py",
        "Category": "Visualization",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Mermaid diagrams, flowcharts, architecture diagrams, auto-generation",
        "Complexity": "Intermediate",
        "Lines": "450+",
        "Integration Value": "Medium - adds visual explanations to responses",
        "Status": "Ready to Clone",
        "Notes": "Generates Mermaid diagrams automatically. Creates flowcharts, architecture diagrams, concept maps."
    },
    
    # HIGH PRIORITY - Semantic Scholar Integration
    {
        "Priority": "HIGH",
        "File": "17_research_assistant_semantic_scholar.py",
        "Category": "Academic Research",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Semantic Scholar API, citation graphs, paper recommendations, research trends",
        "Complexity": "Intermediate",
        "Lines": "400+",
        "Integration Value": "Medium - enhances research capabilities with citation analysis",
        "Status": "Ready to Clone",
        "Notes": "Integrates Semantic Scholar for deep academic research. Shows citation networks and paper recommendations."
    },
    
    # MEDIUM PRIORITY - Rich UI
    {
        "Priority": "MEDIUM",
        "File": "08_research_assistant_rich_ui.py",
        "Category": "User Interface",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Rich library, panels, tables, progress bars, styling",
        "Complexity": "Basic",
        "Lines": "350+",
        "Integration Value": "Low - cosmetic improvements",
        "Status": "Reference Only",
        "Notes": "Beautiful terminal UI with Rich library. Shows professional formatting."
    },
    
    # MEDIUM PRIORITY - Memory Foundations
    {
        "Priority": "MEDIUM",
        "File": "06_research_assistant_with_memory.py",
        "Category": "Memory & Context Management",
        "Learning Value": "⭐⭐⭐",
        "Key Skills": "Basic conversation history, context preservation",
        "Complexity": "Basic",
        "Lines": "300+",
        "Integration Value": "Low - superseded by 09_smart_memory",
        "Status": "Reference Only",
        "Notes": "Basic memory implementation. Study for concepts, but use 09 for production."
    },
    
    # MEDIUM PRIORITY - Accurate Memory
    {
        "Priority": "MEDIUM",
        "File": "07_research_assistant_accurate_memory.py",
        "Category": "Memory & Context Management",
        "Learning Value": "⭐⭐⭐",
        "Key Skills": "Improved context accuracy, conversation threading",
        "Complexity": "Intermediate",
        "Lines": "350+",
        "Integration Value": "Low - superseded by 09_smart_memory",
        "Status": "Reference Only",
        "Notes": "Iterative improvement on memory. Shows progression to smart memory."
    },
    
    # MEDIUM PRIORITY - Expanded Features
    {
        "Priority": "MEDIUM",
        "File": "10_research_assistant_expanded.py",
        "Category": "Feature Expansion",
        "Learning Value": "⭐⭐⭐",
        "Key Skills": "Multiple tools, command handling, feature integration",
        "Complexity": "Intermediate",
        "Lines": "400+",
        "Integration Value": "Low - shows evolutionary step",
        "Status": "Reference Only",
        "Notes": "Mid-stage evolution. Shows how to add features incrementally."
    },
    
    # LOW PRIORITY - Simplest Agent
    {
        "Priority": "LOW",
        "File": "01_simplest_agent.py",
        "Category": "Fundamentals",
        "Learning Value": "⭐⭐",
        "Key Skills": "Basic Anthropic API call, single interaction",
        "Complexity": "Beginner",
        "Lines": "10",
        "Integration Value": "None - too simple",
        "Status": "Reference Only",
        "Notes": "Hello World level. Single API call example for absolute beginners."
    },
    
    # SUPPORTING DOCUMENTATION
    {
        "Priority": "REFERENCE",
        "File": "QUICK_START.md",
        "Category": "Documentation",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Setup instructions, usage examples, troubleshooting",
        "Complexity": "N/A",
        "Lines": "257",
        "Integration Value": "High - essential onboarding guide",
        "Status": "Clone & Adapt",
        "Notes": "Complete setup guide for PDF analysis. Shows excellent documentation structure."
    },
    
    {
        "Priority": "REFERENCE",
        "File": "API_INTEGRATION_SUMMARY.md",
        "Category": "Documentation",
        "Learning Value": "⭐⭐⭐⭐⭐",
        "Key Skills": "API integration patterns, authentication, rate limits, examples",
        "Complexity": "N/A",
        "Lines": "347",
        "Integration Value": "High - shows how to integrate 5 free APIs",
        "Status": "Clone & Study",
        "Notes": "Comprehensive guide to integrating arXiv, Wikipedia, GitHub, Dictionary, Open Library."
    },
    
    {
        "Priority": "REFERENCE",
        "File": "README_API_TOOLS.md",
        "Category": "Documentation",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Tool documentation, API references, usage patterns",
        "Complexity": "N/A",
        "Lines": "Unknown",
        "Integration Value": "Medium - technical reference",
        "Status": "Reference",
        "Notes": "Technical documentation for API tools."
    },
    
    # UTILITY TOOLS
    {
        "Priority": "UTILITY",
        "File": "extract_pdf.py",
        "Category": "Utility Scripts",
        "Learning Value": "⭐⭐⭐",
        "Key Skills": "PDF text extraction, batch processing",
        "Complexity": "Basic",
        "Lines": "150+",
        "Integration Value": "Low - standalone utility",
        "Status": "Keep as Tool",
        "Notes": "Extracts text from PDFs. Useful utility but not agent-specific."
    },
    
    {
        "Priority": "UTILITY",
        "File": "test_token_counting_accuracy.py",
        "Category": "Testing",
        "Learning Value": "⭐⭐⭐",
        "Key Skills": "Token counting validation, accuracy testing",
        "Complexity": "Basic",
        "Lines": "100+",
        "Integration Value": "Medium - validates token tracking",
        "Status": "Reference for Testing",
        "Notes": "Tests accuracy of token counting. Useful for validating 14_token_tracking."
    },
    
    {
        "Priority": "UTILITY",
        "File": "requirements.txt",
        "Category": "Configuration",
        "Learning Value": "⭐⭐⭐⭐",
        "Key Skills": "Dependency management",
        "Complexity": "N/A",
        "Lines": "20+",
        "Integration Value": "Critical - install all dependencies",
        "Status": "Clone Immediately",
        "Notes": "Lists all required packages. Essential for setup."
    }
]

# Create Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Experiments Analysis"

# Headers
headers = [
    "Priority",
    "File",
    "Category",
    "Learning Value",
    "Key Skills",
    "Complexity",
    "Lines of Code",
    "Integration Value",
    "Status",
    "Notes"
]

# Write headers with formatting
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col)
    cell.value = header
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Priority colors
priority_colors = {
    "CRITICAL": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    "HIGH": PatternFill(start_color="FFA94D", end_color="FFA94D", fill_type="solid"),
    "MEDIUM": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
    "LOW": PatternFill(start_color="95E1D3", end_color="95E1D3", fill_type="solid"),
    "REFERENCE": PatternFill(start_color="A8DADC", end_color="A8DADC", fill_type="solid"),
    "UTILITY": PatternFill(start_color="DDA15E", end_color="DDA15E", fill_type="solid"),
}

# Write data
for row_idx, item in enumerate(experiments_analysis, 2):
    ws.cell(row=row_idx, column=1).value = item["Priority"]
    ws.cell(row=row_idx, column=1).fill = priority_colors.get(item["Priority"])
    ws.cell(row=row_idx, column=1).font = Font(bold=True)
    
    ws.cell(row=row_idx, column=2).value = item["File"]
    ws.cell(row=row_idx, column=2).font = Font(name="Consolas")
    
    ws.cell(row=row_idx, column=3).value = item["Category"]
    ws.cell(row=row_idx, column=4).value = item["Learning Value"]
    ws.cell(row=row_idx, column=5).value = item["Key Skills"]
    ws.cell(row=row_idx, column=6).value = item["Complexity"]
    ws.cell(row=row_idx, column=7).value = item["Lines"]
    ws.cell(row=row_idx, column=8).value = item["Integration Value"]
    ws.cell(row=row_idx, column=9).value = item["Status"]
    ws.cell(row=row_idx, column=10).value = item["Notes"]
    
    # Wrap text for all cells
    for col in range(1, 11):
        ws.cell(row=row_idx, column=col).alignment = Alignment(wrap_text=True, vertical="top")

# Adjust column widths
column_widths = [12, 35, 22, 15, 45, 12, 12, 20, 18, 60]
for col, width in enumerate(column_widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = width

# Set row heights
ws.row_dimensions[1].height = 30
for row in range(2, len(experiments_analysis) + 2):
    ws.row_dimensions[row].height = 60

# Create summary sheet
ws_summary = wb.create_sheet("Summary")
ws_summary.column_dimensions['A'].width = 30
ws_summary.column_dimensions['B'].width = 50

summary_data = [
    ["EXPERIMENTS FOLDER ANALYSIS", ""],
    ["", ""],
    ["Total Files Analyzed", len(experiments_analysis)],
    ["", ""],
    ["CRITICAL Priority (Clone First)", sum(1 for x in experiments_analysis if x["Priority"] == "CRITICAL")],
    ["HIGH Priority (Clone Next)", sum(1 for x in experiments_analysis if x["Priority"] == "HIGH")],
    ["MEDIUM Priority (Reference)", sum(1 for x in experiments_analysis if x["Priority"] == "MEDIUM")],
    ["LOW Priority (Optional)", sum(1 for x in experiments_analysis if x["Priority"] == "LOW")],
    ["REFERENCE (Documentation)", sum(1 for x in experiments_analysis if x["Priority"] == "REFERENCE")],
    ["UTILITY (Tools)", sum(1 for x in experiments_analysis if x["Priority"] == "UTILITY")],
    ["", ""],
    ["TOP 3 CRITICAL FILES", ""],
    ["1. 11_multi_agent_research_system.py", "Multi-agent architecture - Essential for Week 5"],
    ["2. 09_research_assistant_smart_memory.py", "Smart memory system - Solves context limits"],
    ["3. 14_research_assistant_token_tracking.py", "Token tracking - Needed for Week 4 metrics"],
    ["", ""],
    ["INTEGRATION RECOMMENDATIONS", ""],
    ["Week 4 (Current)", "Clone: 14_token_tracking for performance metrics"],
    ["Week 5 (Next)", "Clone: 11_multi_agent + 09_smart_memory for multi-agent system"],
    ["Week 6 (Future)", "Clone: 05_apis + 16_obsidian for knowledge base"],
    ["", ""],
    ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")]
]

for row_idx, row_data in enumerate(summary_data, 1):
    ws_summary.cell(row=row_idx, column=1).value = row_data[0]
    ws_summary.cell(row=row_idx, column=2).value = row_data[1]
    
    if row_idx == 1:
        ws_summary.cell(row=row_idx, column=1).font = Font(bold=True, size=14)
        ws_summary.cell(row=row_idx, column=1).fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        ws_summary.cell(row=row_idx, column=1).font = Font(bold=True, size=14, color="FFFFFF")
    elif row_data[0] in ["TOP 3 CRITICAL FILES", "INTEGRATION RECOMMENDATIONS"]:
        ws_summary.cell(row=row_idx, column=1).font = Font(bold=True, size=12)
        ws_summary.cell(row=row_idx, column=1).fill = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")

# Save
filename = "c:\\Users\\Stang3x\\Documents\\Gemini projects\\EXPERIMENTS_ANALYSIS.xlsx"
wb.save(filename)
print(f"✓ Analysis complete: {filename}")
print(f"✓ Total files analyzed: {len(experiments_analysis)}")
print(f"✓ CRITICAL priority: {sum(1 for x in experiments_analysis if x['Priority'] == 'CRITICAL')}")
print(f"✓ HIGH priority: {sum(1 for x in experiments_analysis if x['Priority'] == 'HIGH')}")
